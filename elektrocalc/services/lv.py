from __future__ import annotations

from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

import csv
from openpyxl import load_workbook
from sqlalchemy import select, or_
from sqlalchemy.orm import Session

from elektrocalc.db.models import LvDocument, LvPosition, LvPositionMatch, Item

def _to_qty(v: Any) -> str:
    if v is None or v == "":
        return "1"
    s = str(v).strip().replace(",", ".")
    try:
        d = Decimal(s)
        if d <= 0:
            return "1"
        return str(d.normalize())
    except (InvalidOperation, ValueError):
        return "1"

def _to_unit(v: Any) -> str:
    if v is None or v == "":
        return "pcs"
    s = str(v).strip()
    m = {
        "stk": "pcs", "stück": "pcs", "st": "pcs", "pcs": "pcs",
        "m": "m", "lfm": "m",
        "m2": "m²", "qm": "m²", "m²": "m²",
        "h": "h", "std": "h", "stunden": "h",
    }
    key = s.lower().replace(".", "")
    return m.get(key, s)

def list_lv_documents(s: Session, project_id: int) -> list[LvDocument]:
    return list(s.execute(select(LvDocument).where(LvDocument.project_id==project_id).order_by(LvDocument.id.desc())).scalars().all())

def get_lv_document(s: Session, doc_id: int) -> LvDocument:
    return s.execute(select(LvDocument).where(LvDocument.id==doc_id)).scalar_one()

def list_positions(s: Session, doc_id: int) -> list[LvPosition]:
    return list(s.execute(select(LvPosition).where(LvPosition.lv_document_id==doc_id).order_by(LvPosition.id.asc())).scalars().all())

def import_from_txt(s: Session, project_id: int, name: str, text: str, filename: str | None = None) -> LvDocument:
    doc = LvDocument(project_id=project_id, name=name, source_type="txt", filename=filename, created_at=datetime.utcnow())
    s.add(doc)
    s.flush()

    for line in text.splitlines():
        t = line.strip()
        if not t:
            continue
        parts = [p.strip() for p in t.split(";")]
        if len(parts) >= 4:
            pos_no, qty, unit, short = parts[0], parts[1], parts[2], parts[3]
            long_text = ";".join(parts[4:]).strip() if len(parts) > 4 else None
        else:
            pos_no, qty, unit, short, long_text = None, "1", "pcs", t[:255], None

        s.add(LvPosition(
            lv_document_id=doc.id,
            pos_no=pos_no or None,
            short_text=short[:255],
            long_text=long_text,
            qty=_to_qty(qty),
            unit=_to_unit(unit),
            attributes={},
        ))

    s.commit()
    s.refresh(doc)
    return doc

def import_from_csv(s: Session, project_id: int, name: str, csv_path: Path, mapping: dict[str,str | None], delimiter: str = ";", filename: str | None = None) -> LvDocument:
    doc = LvDocument(project_id=project_id, name=name, source_type="csv", filename=filename or csv_path.name, created_at=datetime.utcnow())
    s.add(doc)
    s.flush()

    with open(csv_path, "r", encoding="utf-8-sig", newline="") as f:
        reader = csv.DictReader(f, delimiter=delimiter)
        for row in reader:
            st = (row.get(mapping.get("short_text") or "") or "").strip()
            if not st:
                continue
            s.add(LvPosition(
                lv_document_id=doc.id,
                pos_no=((row.get(mapping.get("pos_no") or "") or "").strip() or None),
                short_text=st[:255],
                long_text=((row.get(mapping.get("long_text") or "") or "").strip() or None),
                qty=_to_qty(row.get(mapping.get("qty") or "")),
                unit=_to_unit(row.get(mapping.get("unit") or "")),
                attributes={},
            ))
    s.commit()
    s.refresh(doc)
    return doc

def import_from_xlsx(s: Session, project_id: int, name: str, xlsx_path: Path, sheet: str, mapping: dict[str,str | None], header_row: int=1, start_row: int=2, filename: str | None = None) -> LvDocument:
    doc = LvDocument(project_id=project_id, name=name, source_type="xlsx", filename=filename or xlsx_path.name, created_at=datetime.utcnow())
    s.add(doc)
    s.flush()

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        headers = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        header_to_idx = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

        def cell(row, key):
            col = mapping.get(key)
            if not col:
                return None
            idx = header_to_idx.get(col)
            if idx is None:
                return None
            return row[idx]

        for row in ws.iter_rows(min_row=start_row, values_only=True):
            st = cell(row, "short_text")
            if st is None or str(st).strip()=="":
                continue
            s.add(LvPosition(
                lv_document_id=doc.id,
                pos_no=str(cell(row,"pos_no")).strip() if cell(row,"pos_no") not in (None,"") else None,
                short_text=str(st).strip()[:255],
                long_text=str(cell(row,"long_text")).strip() if cell(row,"long_text") not in (None,"") else None,
                qty=_to_qty(cell(row,"qty")),
                unit=_to_unit(cell(row,"unit")),
                attributes={},
            ))
        s.commit()
    finally:
        wb.close()
    s.refresh(doc)
    return doc

def set_match_manual(s: Session, position_id: int, item_id: int | None, note: str | None = None) -> None:
    m = s.execute(select(LvPositionMatch).where(LvPositionMatch.lv_position_id==position_id)).scalar_one_or_none()
    if m is None:
        m = LvPositionMatch(lv_position_id=position_id, item_id=item_id, confidence="manual", note=note, created_at=datetime.utcnow())
        s.add(m)
    else:
        m.item_id = item_id
        m.confidence = "manual"
        m.note = note
    s.commit()

def search_items(s: Session, q: str, limit: int = 20) -> list[Item]:
    qn = (q or "").strip()
    if len(qn) < 2:
        return []
    like = f"%{qn}%"
    return list(s.execute(
        select(Item).where(or_(Item.article_no.like(like), Item.name.like(like), Item.manufacturer.like(like))).order_by(Item.id.desc()).limit(limit)
    ).scalars().all())
