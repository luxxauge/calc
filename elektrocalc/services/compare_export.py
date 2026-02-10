from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

def _autosize(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[col_letter].width = min(60, max(10, max_len + 2))

def build_compare_xlsx(data: dict[str, Any]) -> bytes:
    """Create an XLSX export for LV compare."""
    wb = Workbook()

    # Summary sheet
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Projekt", data["project_name"]])
    ws.append(["LV", data["doc_name"]])
    ws.append([])
    ws.append(["Variante", "Material €", "Lohn €", "Min", "Gesamt €", "Estimate-ID"])
    for s in data["summaries"]:
        ws.append([
            s["variant"],
            round(s["material_cent"]/100, 2),
            round(s["labor_cent"]/100, 2),
            s["labor_minutes"],
            round(s["total_cent"]/100, 2),
            s.get("estimate_id") or "",
        ])
    _autosize(ws)

    # Delta sheet
    ws2 = wb.create_sheet("Delta")
    ws2.append(["Pos", "Kurztext", "Standard €", "Komfort €", "Premium €", "ΔK-Std €", "ΔP-Std €"])
    for r in data["rows"]:
        std = r["std_total"]/100
        kom = r["kom_total"]/100
        pre = r["pre_total"]/100
        ws2.append([r["pos_no"], r["short_text"], round(std,2), round(kom,2), round(pre,2), round(kom-std,2), round(pre-std,2)])
    _autosize(ws2)

    # Extras sheet
    ws3 = wb.create_sheet("Extras")
    ws3.append(["Variante", "Beschreibung", "Rate-Key", "Qty", "Unit", "Material €", "Lohn €", "Total €"])
    for variant, extras in data["extras_by_variant"].items():
        for x in extras:
            ws3.append([
                variant,
                x.description,
                x.rate_key,
                x.qty,
                x.unit,
                round(x.material_total_cent/100, 2),
                round(x.labor_total_cent/100, 2),
                round(x.total_cent/100, 2),
            ])
    _autosize(ws3)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()

def build_compare_pdf(data: dict[str, Any]) -> bytes:
    """Create a compact PDF report for LV compare."""
    bio = BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("Variantenvergleich (LV)", styles["Title"]))
    story.append(Paragraph(f"Projekt: {data['project_name']} – LV: {data['doc_name']}", styles["Normal"]))
    story.append(Spacer(1, 10))

    # Summary table
    sum_table = [["Variante", "Material €", "Lohn €", "Min", "Gesamt €"]]
    for s in data["summaries"]:
        sum_table.append([
            s["variant"],
            f"{s['material_cent']/100:.2f}",
            f"{s['labor_cent']/100:.2f}",
            str(s["labor_minutes"]),
            f"{s['total_cent']/100:.2f}",
        ])
    t = Table(sum_table, hAlign="LEFT")
    t.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0), colors.lightgrey),
        ("GRID",(0,0),(-1,-1), 0.25, colors.grey),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("ALIGN",(1,1),(-1,-1),"RIGHT"),
    ]))
    story.append(Paragraph("Summary", styles["Heading2"]))
    story.append(t)
    story.append(Spacer(1, 10))

    # Delta table (first N rows)
    rows = data["rows"]
    max_rows = 40  # keep PDF readable
    delta_table = [["Pos","Kurztext","Std €","Kom €","Pre €","ΔK","ΔP"]]
    for r in rows[:max_rows]:
        std = r["std_total"]/100
        kom = r["kom_total"]/100
        pre = r["pre_total"]/100
        delta_table.append([r["pos_no"], r["short_text"], f"{std:.2f}", f"{kom:.2f}", f"{pre:.2f}", f"{(kom-std):.2f}", f"{(pre-std):.2f}"])
    t2 = Table(delta_table, colWidths=[40, 320, 70, 70, 70, 60, 60], hAlign="LEFT")
    t2.setStyle(TableStyle([
        ("BACKGROUND",(0,0),(-1,0), colors.lightgrey),
        ("GRID",(0,0),(-1,-1), 0.25, colors.grey),
        ("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),
        ("VALIGN",(0,0),(-1,-1),"TOP"),
        ("ALIGN",(2,1),(-1,-1),"RIGHT"),
    ]))
    story.append(Paragraph("Delta (erste 40 Zeilen)", styles["Heading2"]))
    story.append(t2)

    doc.build(story)
    return bio.getvalue()
