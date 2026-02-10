from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from sqlalchemy.orm import Session

from openpyxl import load_workbook

from elektrocalc.db.models import Item, ItemPrice, ImportBatch
from elektrocalc.util.hashing import sha256_hex

DEFAULT_HEADERS = {
    "article_no": ["artikel", "artikelnummer", "artnr", "nr", "artikel-nr", "id"],
    "name": ["name", "bezeichnung", "artikeltext", "kurztext", "titel"],
    "description": ["beschreibung", "langtext", "text"],
    "manufacturer": ["hersteller", "marke"],
    "category": ["kategorie", "warengruppe", "gruppe"],
    "unit": ["einheit", "meh", "unit"],
    "price_ek": ["ek", "einkauf", "netto ek", "preis ek"],
    "price_vk": ["vk", "verkauf", "netto vk", "preis vk", "listenpreis"],
    "vat_rate": ["mwst", "ust", "vat", "steuersatz"],
}

def _norm_header(s: str) -> str:
    return "".join(ch.lower() for ch in str(s).strip() if ch.isalnum() or ch in (" ", "_", "-")).replace("-", " ").replace("_"," ").strip()

def guess_mapping(headers: list[str]) -> dict[str, str | None]:
    norm = [_norm_header(h) for h in headers]
    mapping: dict[str, str | None] = {k: None for k in DEFAULT_HEADERS.keys()}
    for field, candidates in DEFAULT_HEADERS.items():
        for i, h in enumerate(norm):
            if any(c in h for c in candidates):
                mapping[field] = headers[i]
                break
    return mapping

def list_sheets(xlsx_path: Path) -> list[str]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        return wb.sheetnames
    finally:
        wb.close()

def read_headers(xlsx_path: Path, sheet: str, header_row: int = 1) -> list[str]:
    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    try:
        ws = wb[sheet]
        row = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        headers = [str(c).strip() if c is not None else "" for c in row]
        return headers
    finally:
        wb.close()

def _to_cent(v: Any) -> int:
    if v is None or v == "":
        return 0
    try:
        # handle "12,34" or "12.34"
        s = str(v).strip().replace("€","").replace(" ", "")
        s = s.replace(",", ".")
        d = Decimal(s)
        return int((d * 100).quantize(Decimal("1")))
    except (InvalidOperation, ValueError):
        return 0

def _to_vat(v: Any) -> str:
    if v is None or v == "":
        return "0.19"
    s = str(v).strip().replace("%","").replace(",", ".")
    try:
        d = Decimal(s)
        if d > 1:
            d = d / 100
        # clamp typical VAT rates
        if d < 0 or d > 1:
            return "0.19"
        return str(d)
    except InvalidOperation:
        return "0.19"

def import_items_from_excel(
    s: Session,
    xlsx_path: Path,
    sheet: str,
    mapping: dict[str, str | None],
    header_row: int = 1,
    start_row: int = 2,
    source: str | None = None,
) -> dict[str, int]:
    raw = xlsx_path.read_bytes()
    h = sha256_hex(raw)
    batch = ImportBatch(kind="items_excel", filename=xlsx_path.name, file_hash=h, created_at=datetime.utcnow(), row_count=0, notes=source)
    s.add(batch)
    s.flush()

    wb = load_workbook(xlsx_path, read_only=True, data_only=True)
    created = 0
    updated = 0
    prices = 0
    try:
        ws = wb[sheet]
        headers = list(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True))[0]
        header_to_idx = {str(h).strip(): i for i, h in enumerate(headers) if h is not None}

        def cell(row, key):
            col = mapping.get(key)
            if not col:
                return None
            idx = header_to_idx.get(col)
            if idx is None:
                return None
            return row[idx]

        for row in ws.iter_rows(min_row=start_row, values_only=True):
            art = cell(row, "article_no")
            name = cell(row, "name")
            if art is None or str(art).strip() == "" or name is None or str(name).strip()=="":
                continue

            article_no = str(art).strip()
            manufacturer = cell(row, "manufacturer")
            manufacturer_s = str(manufacturer).strip() if manufacturer not in (None, "") else None

            desc = cell(row, "description")
            cat = cell(row, "category")
            unit = cell(row, "unit")
            vat = cell(row, "vat_rate")
            price_ek = cell(row, "price_ek")
            price_vk = cell(row, "price_vk")

            unit_s = str(unit).strip() if unit not in (None, "") else "pcs"
            vat_s = _to_vat(vat)

            item = s.execute(select(Item).where(Item.article_no == article_no, Item.manufacturer == manufacturer_s)).scalar_one_or_none()
            if item is None:
                item = Item(
                    article_no=article_no,
                    name=str(name).strip(),
                    description=str(desc).strip() if desc not in (None, "") else None,
                    manufacturer=manufacturer_s,
                    category=str(cat).strip() if cat not in (None, "") else None,
                    unit=unit_s,
                    vat_rate=vat_s,
                    created_at=datetime.utcnow(),
                )
                s.add(item)
                try:
                    s.flush()
                except IntegrityError:
                    s.rollback()
                    item = s.execute(select(Item).where(Item.article_no == article_no, Item.manufacturer == manufacturer_s)).scalar_one()
                    updated += 1
                else:
                    created += 1
            else:
                # update minimal fields (keep name/desc if incoming has)
                item.name = str(name).strip()
                if desc not in (None, ""):
                    item.description = str(desc).strip()
                if cat not in (None, ""):
                    item.category = str(cat).strip()
                if unit not in (None, ""):
                    item.unit = unit_s
                item.vat_rate = vat_s
                updated += 1

            ek_cent = _to_cent(price_ek)
            vk_cent = _to_cent(price_vk)
            if ek_cent or vk_cent:
                s.add(ItemPrice(
                    item_id=item.id,
                    price_ek_cent=ek_cent,
                    price_vk_cent=vk_cent,
                    currency="EUR",
                    valid_from=None,
                    source=source,
                    created_at=datetime.utcnow(),
                ))
                prices += 1

            batch.row_count += 1

        s.commit()
    finally:
        wb.close()

    return {"created": created, "updated": updated, "prices": prices, "rows": batch.row_count}
