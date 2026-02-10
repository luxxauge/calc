from __future__ import annotations

from io import BytesIO
from openpyxl import Workbook
from openpyxl.utils import get_column_letter

from reportlab.lib.pagesizes import A4, landscape
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet

def _autosize(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            v = "" if cell.value is None else str(cell.value)
            max_len = max(max_len, len(v))
        ws.column_dimensions[col_letter].width = min(60, max(10, max_len + 2))

def build_quantities_xlsx(plan_name: str, project_name: str, quantities: dict, room_areas: list[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Quantities"
    ws.append(["Projekt", project_name])
    ws.append(["Plan", plan_name])
    ws.append([])

    ws.append(["Punkte-Typ", "Anzahl"])
    for k, v in (quantities.get("points_by_type") or {}).items():
        ws.append([k, v])

    ws.append([])
    ws.append(["Route-Typ", "Länge (m)"])
    for k, v in (quantities.get("routes_by_type_m") or {}).items():
        ws.append([k, v])

    ws.append([])
    ws.append(["Route total (m)", quantities.get("routes_total_m")])
    _autosize(ws)

    ws2 = wb.create_sheet("Rooms")
    ws2.append(["ID", "Name", "Typ", "m²", "px²"])
    for r in room_areas:
        ws2.append([r.get("id"), r.get("name"), r.get("room_type"), r.get("area_m2"), r.get("area_px2")])
    _autosize(ws2)

    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()

def build_quantities_pdf(plan_name: str, project_name: str, quantities: dict, room_areas: list[dict]) -> bytes:
    bio = BytesIO()
    doc = SimpleDocTemplate(bio, pagesize=landscape(A4), leftMargin=18, rightMargin=18, topMargin=18, bottomMargin=18)
    styles = getSampleStyleSheet()
    story = []
    story.append(Paragraph("Plan Quantities Report", styles["Title"]))
    story.append(Paragraph(f"Projekt: {project_name} – Plan: {plan_name}", styles["Normal"]))
    story.append(Spacer(1, 10))

    pts = [["Punkt-Typ", "Anzahl"]]
    for k, v in (quantities.get("points_by_type") or {}).items():
        pts.append([k, str(v)])
    t = Table(pts, hAlign="LEFT")
    t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    story.append(Paragraph("Punkte", styles["Heading2"]))
    story.append(t)
    story.append(Spacer(1, 10))

    routes = [["Route-Typ", "Länge (m)"]]
    for k, v in (quantities.get("routes_by_type_m") or {}).items():
        routes.append([k, str(v)])
    t2 = Table(routes, hAlign="LEFT")
    t2.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    story.append(Paragraph("Routen", styles["Heading2"]))
    story.append(t2)
    story.append(Spacer(1, 10))

    rms = [["ID", "Name", "Typ", "m²"]]
    for r in room_areas[:40]:
        rms.append([str(r.get("id")), r.get("name") or "", r.get("room_type") or "", str(r.get("area_m2") or "—")])
    t3 = Table(rms, colWidths=[40, 220, 120, 60], hAlign="LEFT")
    t3.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.lightgrey), ("GRID", (0, 0), (-1, -1), 0.25, colors.grey), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold")]))
    story.append(Paragraph("Räume (erste 40)", styles["Heading2"]))
    story.append(t3)

    doc.build(story)
    return bio.getvalue()
